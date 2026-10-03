from pathlib import Path
from pypdf import PdfReader
from pptx import Presentation
from .chunking import chunk_text

def ingest_pdf(path, source_id):
    out = []
    for n, page in enumerate(PdfReader(path).pages, 1):
        out += chunk_text(source_id, page.extract_text() or "", f"page {n}")
    return out

def ingest_pptx(path, source_id):
    out = []
    for n, slide in enumerate(Presentation(path).slides, 1):
        texts = [s.text for s in slide.shapes if hasattr(s, "text") and s.text]
        out += chunk_text(source_id, "\n".join(texts), f"slide {n}")
    return out

def ingest_file(path, source_id):
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf":
        return ingest_pdf(path, source_id)
    if suffix == ".pptx":
        return ingest_pptx(path, source_id)
    raise ValueError("Supported formats: PDF, PPTX")
